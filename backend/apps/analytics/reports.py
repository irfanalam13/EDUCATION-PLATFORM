"""PDF (reportlab) and Excel (openpyxl) renderers for institution analytics."""
from __future__ import annotations

import io
from datetime import date

from . import services


def _institution_rows(institution_id: int, days: int) -> tuple[list[tuple[str, str]], dict]:
    data = services.institution_overview(institution_id, days=days)
    m = data["metrics"]
    a = m["assignments"]
    rows = [
        ("Active students", str(m["active_students"])),
        ("Daily active users", str(m["daily_active_users"])),
        ("Monthly active users", str(m["monthly_active_users"])),
        ("Learning hours", str(m["learning_hours"])),
        ("Quiz performance (%)", str(m["quiz_performance"])),
        ("Average mastery (%)", str(m["avg_mastery"])),
        ("7-day attendance (%)", str(m["attendance_rate"])),
        ("Published assignments", str(a["published_assignments"])),
        ("Submissions", str(a["submissions"])),
        ("Assignment completion (%)", str(a["completion_rate"])),
    ]
    return rows, data


def build_pdf(title: str, period_label: str, rows: list[tuple[str, str]]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=title)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(title, styles["Title"]),
        Paragraph(period_label, styles["Normal"]),
        Spacer(1, 18),
    ]
    table = Table([["Metric", "Value"], *rows], colWidths=[320, 150])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buf.getvalue()


def build_xlsx(title: str, period_label: str, rows: list[tuple[str, str]], trend: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = period_label
    ws.append([])
    header_fill = PatternFill("solid", fgColor="1E293B")
    ws.append(["Metric", "Value"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for label, value in rows:
        ws.append([label, value])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16

    if trend:
        ts = wb.create_sheet("Trend")
        ts.append(["Date", "Active users", "Learning minutes", "Avg quiz score"])
        for cell in ts[1]:
            cell.font = Font(bold=True)
        for point in trend:
            ts.append(
                [
                    point["date"],
                    point["active_users"],
                    point["learning_minutes"],
                    point["avg_quiz_score"],
                ]
            )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_institution_report(institution_id: int, fmt: str, start: date, end: date, title: str) -> bytes:
    days = max((end - start).days + 1, 1)
    rows, data = _institution_rows(institution_id, days)
    period_label = f"Period: {start.isoformat()} to {end.isoformat()}"
    if fmt == "XLSX":
        return build_xlsx(title, period_label, rows, data.get("trend", []))
    return build_pdf(title, period_label, rows)
